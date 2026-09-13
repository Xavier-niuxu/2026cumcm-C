# -*- coding: utf-8 -*-
"""Shared Excel styling for result subtables."""

from openpyxl.styles import Border, Side


THIN = Side(style="thin")


def _make_border(top: bool = False, bottom: bool = False) -> Border:
    """Create a border with thin vertical rules and optional horizontal rules."""
    return Border(
        left=THIN,
        right=THIN,
        top=THIN if top else None,
        bottom=THIN if bottom else None,
    )


def _date_boundary_rows(ws, date_column: int) -> set:
    """Return rows that finish a date group, plus header/last table rows."""
    last_row = ws.max_row
    if last_row < 1:
        return set()

    boundaries = {1, last_row}
    previous_date = None
    for row in range(2, last_row + 1):
        value = ws.cell(row, date_column).value
        if value is None or (isinstance(value, str) and not value.strip()):
            continue
        if previous_date is not None and str(value) != str(previous_date):
            boundaries.add(row - 1)
        previous_date = value
    return boundaries


def style_day_subtable(ws, date_column: int = 1) -> None:
    """Add vertical column rules, an outer border, and horizontal day rules."""
    if ws.max_row < 1:
        return

    ws.sheet_view.showGridLines = True
    boundaries = _date_boundary_rows(ws, date_column)
    for row in range(1, ws.max_row + 1):
        border = _make_border(
            top=row == 1,
            bottom=row in boundaries,
        )
        for col in range(1, ws.max_column + 1):
            ws.cell(row, col).border = border

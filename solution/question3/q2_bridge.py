# -*- coding: utf-8 -*-
"""替代参考解中缺失的 ``forecast_benchmark``：提供 ``load_data()`` 与 ``fourier()``。

参考解的 rolling_q3.py 从 ``../q2/forecast_benchmark.py`` 导入 ``load_data`` 和
``fourier``，但该模块不在交付包里。这里用本仓库问题 2 的傅里叶模型补上：

* ``load_data()``        -> (小区负载 L (365,144), 实际光伏 P (365,144), 电价 p)
* ``fourier(L, d)``      -> 第 d 天的日前负荷预报（144 维），只用 d 之前的历史拟合

问题 2 的模型建模的是净负荷 N = L - P，把光伏序列置 0 即退化为对负荷建模，
因此这里直接复用 ``pred.fourier.FourierPredictor``，不重复实现。
"""

from __future__ import annotations

import sys
from pathlib import Path

import numpy as np
from openpyxl import load_workbook

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
ATT = ROOT / "附件"
Q2_DIR = HERE.parent / "question2"
if str(Q2_DIR) not in sys.path:
    sys.path.append(str(Q2_DIR))

from pred.fourier import FourierPredictor  # noqa: E402

_MODEL = None
_DATES = None


def load_data():
    """返回 (负荷 L, 光伏 P, 电价 p)。"""
    w = load_workbook(ATT / "附件2.xlsx", read_only=True, data_only=True)
    L = np.array([
        [float(x) for x in r[1:]]
        for r in w.worksheets[0].iter_rows(min_row=2, values_only=True)
    ])
    P = np.array([
        [float(x) for x in r[1:]]
        for r in w.worksheets[1].iter_rows(min_row=2, values_only=True)
    ])
    w = load_workbook(ATT / "附件1.xlsx", read_only=True, data_only=True)
    p = np.array([float(r[1]) for r in w.active.iter_rows(min_row=2, values_only=True)])
    return L, P, p


def _build(L):
    """惰性构建问题 2 的负荷预报模型（滚动扩窗，只用历史）。"""
    global _MODEL, _DATES
    if _MODEL is None:
        w = load_workbook(ATT / "附件2.xlsx", read_only=True, data_only=True)
        dates = np.array([r[0] for r in w.worksheets[0].iter_rows(min_row=2, values_only=True)])
        _DATES = dates
        # 把光伏置 0 -> 被解释变量 L - 0 = L
        _MODEL = FourierPredictor(dates, L, dates, np.zeros_like(L))
    return _MODEL


def fourier(L, d: int) -> np.ndarray:
    """第 d 天的日前负荷预报（144 维），只用 d 之前的数据。"""
    model = _build(L)
    date_str = str(_DATES[d])[:10]
    load_pred, _ = model.predict(date_str)
    return np.asarray(load_pred, dtype=float)

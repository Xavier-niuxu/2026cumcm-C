# -*- coding: utf-8 -*-
"""附件 3 整点光伏预报的分析与降尺度（移植自参考解的 forecast_analysis.py）。

* ``load_all()``      : 读附件 2 实际光伏 P（365×144）与附件 3 预报 F（365×4×24）
* ``downscale()``     : 整点预报 → 10 分钟曲线，支持 step / linear / cubic / pchip /
                        historical_shape / linear_mean_preserving
* ``historical_shape``: 用**过去 28 天**的小时内形状比例构造日型（严格因果）
* ``hourly_metrics`` / ``downscale_metrics``: 预报精度与降尺度方法比较

与参考解的唯一差别是目录定位（本项目 q3 在 solution/ 下）以及输出目录。
"""

from __future__ import annotations

from pathlib import Path
import csv
import json

import numpy as np
from openpyxl import load_workbook
from scipy.interpolate import CubicSpline, PchipInterpolator

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]                      # 项目根目录
ATT = ROOT / "附件"
OUT = HERE / "output"
OUT.mkdir(exist_ok=True)
HOURS = [0, 6, 12, 18]


def load_all():
    """返回 (实际光伏 P (365,144), 预报 F (365,4,24))。"""
    w = load_workbook(ATT / "附件2.xlsx", read_only=True, data_only=True)
    P = np.array([
        [float(x) for x in r[1:]]
        for r in w.worksheets[1].iter_rows(min_row=2, values_only=True)
    ])
    w = load_workbook(ATT / "附件3.xlsx", read_only=True, data_only=True)
    rows = list(w.active.iter_rows(min_row=2, values_only=True))
    F = np.zeros((365, 4, 24))
    for d in range(365):
        for j in range(4):
            F[d, j] = np.array(rows[d * 4 + j][2:26], float)
    return P, F


def actual_hour(P, d, h):
    day = d + h // 24
    hour = h % 24
    if day >= len(P):
        return None
    return P[day, hour * 6:(hour + 1) * 6].mean()


def hourly_metrics(P, F):
    rows = []
    errs = {j: [] for j in range(4)}
    for d in range(365):
        for j, issue in enumerate(HOURS):
            for k in range(24):
                a = actual_hour(P, d, issue + k)
                if a is None:
                    continue
                e = F[d, j, k] - a
                errs[j].append((k + 1, e, a, F[d, j, k]))
    for j, issue in enumerate(HOURS):
        A = np.array(errs[j])
        for group, mask in [
            ("all", np.ones(len(A), bool)),
            ("daytime", (A[:, 2] > 1) | (A[:, 3] > 1)),
            ("nighttime", (A[:, 2] <= 1) & (A[:, 3] <= 1)),
        ]:
            x = A[mask, 1]
            rows.append({
                "issue_hour": issue, "group": group,
                "MAE_kW": np.mean(abs(x)), "RMSE_kW": np.sqrt(np.mean(x * x)),
                "bias_forecast_minus_actual_kW": np.mean(x), "n": len(x),
            })
        for k in range(1, 25):
            x = A[A[:, 0] == k, 1]
            rows.append({
                "issue_hour": issue, "group": f"horizon_{k}",
                "MAE_kW": np.mean(abs(x)), "RMSE_kW": np.sqrt(np.mean(x * x)),
                "bias_forecast_minus_actual_kW": np.mean(x), "n": len(x),
            })
    for season, ds in [
        ("winter", list(range(0, 59)) + list(range(334, 365))),
        ("spring", range(59, 151)), ("summer", range(151, 243)),
        ("autumn", range(243, 334)),
    ]:
        ee = []
        for d in ds:
            for j, issue in enumerate(HOURS):
                for k in range(24):
                    a = actual_hour(P, d, issue + k)
                    if a is not None:
                        ee.append(F[d, j, k] - a)
        x = np.array(ee)
        rows.append({
            "issue_hour": "all", "group": season,
            "MAE_kW": np.mean(abs(x)), "RMSE_kW": np.sqrt(np.mean(x * x)),
            "bias_forecast_minus_actual_kW": np.mean(x), "n": len(x),
        })
    return rows


def downscale(vals, method, shape=None):
    """整点序列 → 10 分钟序列（每小时 6 点）。"""
    vals = np.asarray(vals, float)
    x = np.arange(len(vals)) + .5
    xx = np.arange(len(vals) * 6) / 6 + 1 / 12
    if method == "step":
        y = np.repeat(vals, 6)
    elif method == "linear":
        y = np.interp(xx, x, vals, left=vals[0], right=vals[-1])
    elif method == "cubic":
        y = CubicSpline(x, vals, bc_type="natural", extrapolate=True)(xx)
    elif method == "pchip":
        y = PchipInterpolator(x, vals, extrapolate=True)(xx)
    elif method == "historical_shape":
        y = (vals[:, None] * (shape if shape is not None else np.ones((len(vals), 6)))).ravel()
    elif method == "linear_mean_preserving":
        y = np.maximum(np.interp(xx, x, vals, left=vals[0], right=vals[-1]), 0).reshape(-1, 6)
        den = y.mean(1)
        y = np.where(den[:, None] > 1e-12,
                     y * (vals / np.maximum(den, 1e-12))[:, None], vals[:, None])
        y = y.ravel()
    else:
        raise KeyError(method)
    return y


def historical_shape(P, d, start_hour, length=24):
    """过去 28 天的"小时内形状"比例（每小时均值归一），严格因果。"""
    hist = P[max(0, d - 28):d]
    out = np.ones((length, 6))
    for k in range(length):
        hour = (start_hour + k) % 24
        v = hist[:, hour * 6:(hour + 1) * 6]
        den = v.mean(1)
        ok = den > 1
        if ok.any():
            out[k] = np.mean(v[ok] / den[ok, None], axis=0)
    return out


def downscale_metrics(P, F):
    rows = []
    methods = ["step", "linear", "cubic", "pchip", "historical_shape",
               "linear_mean_preserving"]
    for m in methods:
        err, neg, over, n, sun = [], 0, 0, 0, []
        for d in range(31, 365):
            for j, issue in enumerate(HOURS):
                nh = 24 - issue
                shape = historical_shape(P, d, issue, nh)
                pred = downscale(F[d, j, :nh], m, shape)
                act = P[d, issue * 6:]
                raw = pred.copy()
                neg += np.sum(raw < 0)
                over += np.sum(raw > max(F[d, j, :nh].max() * 1.15, 1))
                n += len(raw)
                if m in ("cubic", "pchip", "linear", "linear_mean_preserving"):
                    pred = np.maximum(pred, 0)
                err.extend(pred - act)
                mask = ((act > 0) & (np.r_[0, act[:-1]] <= 1)) | \
                       ((act <= 1) & (np.r_[act[1:], 0] > 0))
                sun.extend((pred - act)[mask])
        e = np.array(err)
        se = np.array(sun)
        rows.append({
            "method": m, "MAE_kW": np.mean(abs(e)), "RMSE_kW": np.sqrt(np.mean(e * e)),
            "bias_kW": np.mean(e), "negative_raw_count": int(neg),
            "overshoot_raw_count": int(over),
            "sunrise_sunset_MAE_kW": np.mean(abs(se)) if len(se) else 0, "n": n,
        })
    return rows


def main():
    P, F = load_all()
    hr = hourly_metrics(P, F)
    dr = downscale_metrics(P, F)
    for name, rows in [("forecast_hourly_metrics.csv", hr),
                       ("downscale_metrics.csv", dr)]:
        with open(OUT / name, "w", newline="", encoding="utf-8-sig") as f:
            w = csv.DictWriter(f, fieldnames=list(rows[0]))
            w.writeheader()
            w.writerows(rows)
    np.savez_compressed(OUT / "pv_forecasts.npz", pv=P, forecast=F)
    print(json.dumps({
        "issue_all": [x for x in hr if x["group"] == "all"],
        "season": [x for x in hr if x["issue_hour"] == "all"],
        "downscale": dr,
    }, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()

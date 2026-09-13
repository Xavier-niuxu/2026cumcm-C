#!/usr/bin/env python3
from __future__ import annotations
from pathlib import Path
from time import perf_counter
import csv,json,sys
import numpy as np
from scipy.optimize import linprog
from scipy.stats import norm,rankdata
from openpyxl import load_workbook

HERE=Path(__file__).resolve().parent;ROOT=HERE.parents[1];ATT=ROOT/'附件';OUT=HERE/'output';OUT.mkdir(exist_ok=True)
# 参考解从 ../q2/forecast_benchmark 与 ../q3/forecast_analysis 取负荷预报与 PV 降尺度；
# 这两个模块不在交付包里，这里用本地等价实现（q2_bridge / pv_tools）替代。
sys.path.insert(0,str(HERE));from q2_bridge import load_data,fourier  # noqa: E402
from pv_tools import load_all,downscale,historical_shape,HOURS  # noqa: E402
DT=1/6;ETA=.9;LIM=5000*DT;EMIN=1200.;EMAX=10800.;TARGET=6000.;START=31;N=144

def kmeans(X,k,seed=2026,iters=40):
 """Small deterministic NumPy k-means, avoiding external OpenMP runtime conflicts."""
 X=np.asarray(X,float);rng=np.random.default_rng(seed);cent=X[rng.choice(len(X),k,replace=False)].copy();lab=np.zeros(len(X),int)
 for _ in range(iters):
  new=np.argmin(((X[:,None,:]-cent[None,:,:])**2).sum(2),axis=1)
  if np.array_equal(new,lab):break
  lab=new
  for j in range(k):
   if np.any(lab==j):cent[j]=X[lab==j].mean(0)
 return cent,lab

def load_prices_dates():
 w=load_workbook(ATT/'附件4.xlsx',read_only=True,data_only=True);rows=list(w.active.iter_rows(min_row=2,values_only=True));return np.array([[float(x) for x in r[1:]] for r in rows]),[r[0] for r in rows]
def labels():
 out=[]
 for t in range(N):
  a=t*10;b=(t+1)*10;f=lambda x:'0:00+1' if x==1440 else f'{x//60}:{x%60:02d}';out.append(f'{f(a)}-{f(b)}')
 return out
def event_count(e):return sum(e[i]>5e-5 and (i==0 or e[i-1]<=5e-5) for i in range(len(e)))

def reduce_joint(en,ep,k=5,mode='joint_block',seed=2026):
 """Return net-error paths, price-error paths, probabilities; all inputs are completed past days."""
 m=len(en)
 if m==0:return np.zeros((1,en.shape[1] if en.ndim==2 else N)),np.zeros((1,ep.shape[1] if ep.ndim==2 else N)),np.ones(1)
 en=en[-90:];ep=ep[-90:];m=len(en);k=min(k,m)
 sn=en.std()+1e-9;sp=ep.std()+1e-9
 if mode in ('joint_block','cluster_kmeans','conditional_regime'):
  Z=np.c_[en/sn,ep/sp];c,lab=kmeans(Z,k,seed);w=np.bincount(lab,minlength=k)/m;return c[:,:en.shape[1]]*sn,c[:,en.shape[1]:]*sp,w
 if mode=='independent':
  ca,_=kmeans(en/sn,k,seed);cb,_=kmeans(ep/sp,k,seed+7);return ca*sn,np.roll(cb,1,axis=0)*sp,np.ones(k)/k
 rng=np.random.default_rng(seed+m)
 if mode=='multivariate_gaussian':
  # Low-rank Gaussian paths retain within-variable temporal covariance and cross covariance.
  Z=np.c_[en,ep];mu=Z.mean(0);U,s,V=np.linalg.svd(Z-mu,full_matrices=False);coef=rng.normal(size=(max(30,k*6),len(s)))*s/np.sqrt(max(m-1,1));sim=mu+coef@V
 elif mode in ('gaussian_copula','vine_copula'):
  # Empirical marginal paths coupled by Gaussian ranks; vine uses tail-severity ordering as a second conditioning edge.
  an=en.mean(1);ap=ep.mean(1);zn=norm.ppf((rankdata(an)-.5)/m);zp=norm.ppf((rankdata(ap)-.5)/m);rho=np.corrcoef(zn,zp)[0,1]
  z1=rng.normal(size=max(30,k*6));z2=rho*z1+np.sqrt(max(1-rho*rho,1e-6))*rng.normal(size=len(z1));
  ix=np.clip((norm.cdf(z1)*m).astype(int),0,m-1);iy=np.clip((norm.cdf(z2)*m).astype(int),0,m-1)
  if mode=='vine_copula':
   tail=np.argsort(np.maximum(en,0).sum(1));iy=np.where(z1>1,tail[np.clip((norm.cdf(z2)*m).astype(int),0,m-1)],iy)
  sim=np.c_[en[ix],ep[iy]]
 else:raise KeyError(mode)
 Z=np.c_[sim[:,:en.shape[1]]/sn,sim[:,en.shape[1]:]/sp];c,lab=kmeans(Z,k,seed);w=np.bincount(lab,minlength=k)/len(Z);return c[:,:en.shape[1]]*sn,c[:,en.shape[1]:]*sp,w

def stochastic_commit(point,pricehat,s0,en,ep,mode='joint_block',risk_lambda=0,alpha=.95,tail_weight=False):
 ne,pe,w=reduce_joint(en,ep,5,mode);K=len(w);n=len(point);ps=np.maximum(pricehat[None,:]+pe,.01)
 if tail_weight:
  sev=np.maximum(ne,0).sum(1)*ps.mean(1);w=w*(1+sev/(sev.mean()+1e-9));w/=w.sum()
 base=n+K*4*n;extra=(1+K if risk_lambda>0 else 0);nv=base+extra;o=np.zeros(nv)
 o[:n]=(w[:,None]*ps).sum(0)
 for k in range(K):o[n+k*4*n+3*n:n+(k+1)*4*n]=5*w[k]*ps[k]
 if risk_lambda>0:o[base]=risk_lambda;o[base+1:]=risk_lambda*w/(1-alpha)
 Ae=[];be=[];Au=[];bu=[]
 for k in range(K):
  off=n+k*4*n;net=point+ne[k]
  for t in range(n):
   r=np.zeros(nv);r[off+t]=-ETA;r[off+n+t]=1/ETA;r[off+2*n+t]=1
   if t:r[off+2*n+t-1]=-1;rhs=0
   else:rhs=s0
   Ae.append(r);be.append(rhs)
   r=np.zeros(nv);r[t]=-1;r[off+t]=1;r[off+n+t]=-1;r[off+3*n+t]=-1;Au.append(r);bu.append(-net[t]*DT)
  r=np.zeros(nv);r[off+3*n-1]=1;Ae.append(r);be.append(TARGET)
  if risk_lambda>0:
   # scenario cost - VaR - excess <= 0
   r=np.zeros(nv);r[:n]=ps[k];r[off+3*n:off+4*n]=5*ps[k];r[base]=-1;r[base+1+k]=-1;Au.append(r);bu.append(0)
 bds=[(0,None)]*n
 for _ in range(K):bds += [(0,LIM)]*(2*n)+[(EMIN,EMAX)]*n+[(0,None)]*n
 if risk_lambda>0:bds += [(None,None)]+[(0,None)]*K
 z=linprog(o,A_ub=np.array(Au),b_ub=np.array(bu),A_eq=np.array(Ae),b_eq=np.array(be),bounds=bds,method='highs')
 if not z.success:raise RuntimeError(z.message)
 return z.x[:n]

def execute(g,actual,s0):
 c=np.zeros(N);d=np.zeros(N);e=np.zeros(N);S=np.empty(N+1);S[0]=s0
 for t in range(N):
  gap=actual[t]*DT-g[t]
  if gap>0:d[t]=min(gap,LIM,(S[t]-EMIN)*ETA)
  else:c[t]=min(-gap,LIM,(EMAX-S[t])/ETA)
  S[t+1]=S[t]+ETA*c[t]-d[t]/ETA;e[t]=max(0,gap+c[t]-d[t])
 return c,d,S,e

def q42_backtest(name,net,net_hat,prices,price_hat,mode='joint_block',risk_lambda=0,alpha=.95,tail=False,continuous=True):
 tic=perf_counter();days=[];s0=TARGET
 for i in range(len(net)):
  en=net[:i]-net_hat[:i];ep=prices[:i]-price_hat[:i]
  if name=='deterministic':en=np.empty((0,N));ep=np.empty((0,N))
  if name=='robust' and len(en):
   # upper-tail deterministic shortage/price path
   en=np.quantile(en[-90:],.95,axis=0)[None,:];ep=np.quantile(ep[-90:],.90,axis=0)[None,:]
  g=stochastic_commit(net_hat[i],price_hat[i],s0,en,ep,mode,risk_lambda,alpha,tail)
  c,d,S,e=execute(g,net[i],s0);pc=float(g@prices[i]);ec=float(e@(5*prices[i]));cost=pc+ec
  days.append(dict(g=g,c=c,d=d,soc=S,e=e,plan_cost=pc,emergency_cost=ec,total_cost=cost))
  if continuous:s0=float(S[-1])
 out={'problem':'Q4-2','name':name,'scenario':mode,'risk_lambda':risk_lambda,'alpha':alpha,'plan_cost':sum(x['plan_cost'] for x in days),'emergency_cost':sum(x['emergency_cost'] for x in days),'total_cost':sum(x['total_cost'] for x in days),'emergency_kwh':sum(x['e'].sum() for x in days),'events':sum(event_count(x['e']) for x in days),'throughput':sum(x['c'].sum()+x['d'].sum() for x in days),'terminal_soc':days[-1]['soc'][-1],'violations':sum(x['soc'].min()<EMIN-1e-6 or x['soc'].max()>EMAX+1e-6 for x in days),'daily_q95':float(np.quantile([x['total_cost'] for x in days],.95)),'daily_max':max(x['total_cost'] for x in days),'runtime_s':perf_counter()-tic}
 return out,days

def pv_curve(P,F,d,j):
 issue=HOURS[j];return np.maximum(downscale(F[d,j,:24-issue],'linear',historical_shape(P,d,issue,24-issue)),0)
def price_update(base,actual,d,h):
 """At h, update future price only with prices observed strictly before h."""
 start=h*6;x=base[d,start:].copy()
 if start:
  obs=actual[d,:start]-base[d,:start];level=np.mean(obs[-min(12,start):]);decay=np.exp(-np.arange(len(x))/36);x+=level*decay
 return np.maximum(x,.01)
def q43_histories(L,P,F,Lhat,prices,phat,d,pvj,h,window=90):
 start=h*6;en=[];ep=[]
 for r in range(max(START,d-window),d):
  pvfull=pv_curve(P,F,r,pvj);offset=start-HOURS[pvj]*6
  point=Lhat[r,start:]-pvfull[offset:];en.append((L[r,start:]-P[r,start:])-point)
  pu=price_update(phat,prices,r,h);ep.append(prices[r,start:]-pu)
 if not en:return np.empty((0,N-start)),np.empty((0,N-start))
 return np.asarray(en),np.asarray(ep)

def rolling_adjust(point,pricehat,s0,prev,en,ep,mode='joint_block',risk_lambda=0,alpha=.95,tail=False):
 ne,pe,w=reduce_joint(en,ep,5,mode);K=len(w);n=len(point);ps=np.maximum(pricehat[None]+pe,.01);adjust=prev is not None
 qbase=n+(2*n if adjust else 0);extra=(1+K if risk_lambda>0 else 0);nv=qbase+K*4*n+extra;o=np.zeros(nv)
 if adjust:
  o[n:2*n]=1.5*(w[:,None]*ps).sum(0);o[2*n:3*n]=-.5*(w[:,None]*ps).sum(0)
 else:o[:n]=(w[:,None]*ps).sum(0)
 for k in range(K):o[qbase+k*4*n+3*n:qbase+(k+1)*4*n]=5*w[k]*ps[k]
 riskbase=qbase+K*4*n
 if risk_lambda>0:o[riskbase]=risk_lambda;o[riskbase+1:]=risk_lambda*w/(1-alpha)
 Ae=[];be=[];Au=[];bu=[]
 if adjust:
  for t in range(n):r=np.zeros(nv);r[t]=1;r[n+t]=-1;r[2*n+t]=1;Ae.append(r);be.append(prev[t])
 for k in range(K):
  off=qbase+k*4*n;net=point+ne[k]
  for t in range(n):
   r=np.zeros(nv);r[off+t]=-ETA;r[off+n+t]=1/ETA;r[off+2*n+t]=1
   if t:r[off+2*n+t-1]=-1;rhs=0
   else:rhs=s0
   Ae.append(r);be.append(rhs)
   r=np.zeros(nv);r[t]=-1;r[off+t]=1;r[off+n+t]=-1;r[off+3*n+t]=-1;Au.append(r);bu.append(-net[t]*DT)
  r=np.zeros(nv);r[off+3*n-1]=1;Ae.append(r);be.append(TARGET)
  if risk_lambda>0:
   r=np.zeros(nv);r[off+3*n:off+4*n]=5*ps[k]
   if adjust:r[n:2*n]=1.5*ps[k];r[2*n:3*n]=-.5*ps[k]
   else:r[:n]=ps[k]
   r[riskbase]=-1;r[riskbase+1+k]=-1;Au.append(r);bu.append(0)
 bds=[(0,None)]*n+([(0,None)]*(2*n) if adjust else [])
 for _ in range(K):bds += [(0,LIM)]*(2*n)+[(EMIN,EMAX)]*n+[(0,None)]*n
 if risk_lambda>0:bds += [(None,None)]+[(0,None)]*K
 z=linprog(o,A_ub=np.array(Au),b_ub=np.array(bu),A_eq=np.array(Ae),b_eq=np.array(be),bounds=bds,method='highs')
 if not z.success:raise RuntimeError(z.message)
 return z.x[:n]

def execute_stage(q,actual,start,end,S):
 c=np.zeros(end-start);d=np.zeros(end-start);e=np.zeros(end-start);soc=np.empty(end-start+1);soc[0]=S
 for z,t in enumerate(range(start,end)):
  gap=actual[t]*DT-q[t]
  if gap>0:d[z]=min(gap,LIM,(soc[z]-EMIN)*ETA)
  else:c[z]=min(-gap,LIM,(EMAX-soc[z])/ETA)
  soc[z+1]=soc[z]+ETA*c[z]-d[z]/ETA;e[z]=max(0,gap+c[z]-d[z])
 return c,d,e,soc

def q43_day(d,L,P,F,Lhat,prices,phat,s0,mode='joint_block',risk_lambda=0,alpha=.95,deterministic=False,pv_updates=True,price_updates=True):
 actual=L[d]-P[d];current=None;g0=None;gfinal=np.zeros(N);c=np.zeros(N);dis=np.zeros(N);e=np.zeros(N);soc=np.empty(N+1);soc[0]=s0;adjust_cost=0;adjust_kwh=0
 for j,h in enumerate(HOURS):
  start=h*6;end=(HOURS[j+1]*6 if j+1<len(HOURS) else N)
  jj=j if pv_updates else 0;point=Lhat[d,start:]-pv_curve(P,F,d,jj)[start-HOURS[jj]*6:]
  pnow=price_update(phat,prices,d,h) if price_updates else phat[d,start:]
  en,ep=q43_histories(L,P,F,Lhat,prices,phat,d,jj,h)
  if deterministic:en=np.empty((0,len(point)));ep=np.empty((0,len(point)))
  prev=None if h==0 else current[start:]
  q=rolling_adjust(point,pnow,soc[start],prev,en,ep,mode,risk_lambda,alpha)
  if h==0:g0=q.copy();current=q.copy()
  else:
   up=np.maximum(q-prev,0);dn=np.maximum(prev-q,0);adjust_cost+=float(1.5*up@prices[d,start:]-.5*dn@prices[d,start:]);adjust_kwh+=float(up.sum()+dn.sum());current[start:]=q
  gfinal[start:end]=current[start:end];cc,dd,ee,ss=execute_stage(current,actual,start,end,soc[start]);c[start:end]=cc;dis[start:end]=dd;e[start:end]=ee;soc[start:end+1]=ss
 plan_cost=float(g0@prices[d]);emergency_cost=float(e@(5*prices[d]));return dict(g0=g0,gfinal=gfinal,c=c,d=dis,soc=soc,e=e,plan_cost=plan_cost,adjustment_cost=adjust_cost,emergency_cost=emergency_cost,total_cost=plan_cost+adjust_cost+emergency_cost,adjustment_kwh=adjust_kwh)

def q43_backtest(name,L,P,F,Lhat,prices,phat,mode='joint_block',risk_lambda=0,alpha=.95,deterministic=False,pv_updates=True,price_updates=True):
 tic=perf_counter();days=[];s0=TARGET
 for day in range(START,365):
  r=q43_day(day,L,P,F,Lhat,prices,phat,s0,mode,risk_lambda,alpha,deterministic,pv_updates,price_updates);days.append(r);s0=float(r['soc'][-1])
 out={'problem':'Q4-3','name':name,'scenario':mode,'risk_lambda':risk_lambda,'alpha':alpha,'pv_updates':pv_updates,'price_updates':price_updates,'plan_cost':sum(x['plan_cost'] for x in days),'adjustment_cost':sum(x['adjustment_cost'] for x in days),'emergency_cost':sum(x['emergency_cost'] for x in days),'total_cost':sum(x['total_cost'] for x in days),'emergency_kwh':sum(x['e'].sum() for x in days),'adjustment_kwh':sum(x['adjustment_kwh'] for x in days),'events':sum(event_count(x['e']) for x in days),'throughput':sum(x['c'].sum()+x['d'].sum() for x in days),'terminal_soc':days[-1]['soc'][-1],'violations':sum(x['soc'].min()<EMIN-1e-6 or x['soc'].max()>EMAX+1e-6 for x in days),'daily_q95':float(np.quantile([x['total_cost'] for x in days],.95)),'daily_max':max(x['total_cost'] for x in days),'runtime_s':perf_counter()-tic}
 return out,days

def save_result42(dates,days):
 wb=load_workbook(ATT/'附件5/result4-2.xlsx');labs=labels();ws=wb['计划购电量']
 for t,x in enumerate(labs,2):ws.cell(1,t,x)
 for i,(date,r) in enumerate(zip(dates[START:],days),2):
  ws.cell(i,1,date)
  for t,x in enumerate(r['g'],2):ws.cell(i,t,round(float(x),4))
  ws.cell(i,146,round(float(r['g'].sum()),4));ws.cell(i,147,round(float(r['total_cost']),4))
 fill_storage_emergency(wb,dates,days,'e');wb.save(OUT/'result4-2.xlsx')
def save_result43(dates,days):
 wb=load_workbook(ATT/'附件5/result4-3.xlsx');labs=labels()
 for sn,key in [('计划购电量','g0'),('调整购电量','gfinal')]:
  ws=wb[sn]
  for t,x in enumerate(labs,2):ws.cell(1,t,x)
  for i,(date,r) in enumerate(zip(dates[START:],days),2):
   ws.cell(i,1,date)
   for t,x in enumerate(r[key],2):ws.cell(i,t,round(float(x),4))
   ws.cell(i,146,round(float(r[key].sum()),4));ws.cell(i,147,round(float(r['total_cost']),4))
 fill_storage_emergency(wb,dates,days,'e');wb.save(OUT/'result4-3.xlsx')
def fill_storage_emergency(wb,dates,days,ekey):
 spans=['0:00-4:00','4:00-8:00','8:00-12:00','12:00-16:00','16:00-20:00','20:00-24:00'];ws=wb['充放电量'];ws.delete_rows(2,ws.max_row);row=2
 for date,r in zip(dates[START:],days):
  for k,sp in enumerate(spans):
   ws.cell(row,1,date if k==0 else None);ws.cell(row,2,sp);ws.cell(row,3,round(float(r['c'][k*24:(k+1)*24].sum()),4));ws.cell(row,4,round(float(r['d'][k*24:(k+1)*24].sum()),4))
   if k==0:ws.cell(row,5,'0:00');ws.cell(row,6,round(float(r['soc'][0]),4))
   if k==1:ws.cell(row,5,'24:00');ws.cell(row,6,round(float(r['soc'][-1]),4))
   row+=1
 ws=wb['紧急购电量'];ws.delete_rows(2,ws.max_row);row=2;labs=labels()
 for date,r in zip(dates[START:],days):
  e=r[ekey];i=0;first=True
  while i<N:
   if e[i]<=5e-5:i+=1;continue
   j=i+1
   while j<N and e[j]>5e-5:j+=1
   ws.cell(row,1,date if first else None);ws.cell(row,2,labs[i].split('-')[0]+'-'+labs[j-1].split('-')[1]);ws.cell(row,3,round(float(e[i:j].sum()),4));row+=1;first=False;i=j
  if first:ws.cell(row,1,date);ws.cell(row,2,'无');ws.cell(row,3,0);row+=1

def main():
 L,P,netall=load_data();_,F=load_all();prices_all,dates=load_prices_dates();pz=np.load(OUT/'price_forecasts.npz');price_hat=pz['HistGB_XGB_fallback'];prices=prices_all[START:]
 Lhat=np.array([fourier(L,d) if d>=7 else L[:max(d,1)].mean(0) for d in range(365)]);net_hat=np.array([fourier(netall,d) for d in range(START,365)]);net=netall[START:]
 rows=[];cache={}
 # Q4-2 scenario dependence and method comparison.
 for mode in ['independent','joint_block','multivariate_gaussian','gaussian_copula','vine_copula','conditional_regime']:
  r,ds=q42_backtest('joint_structure',net,net_hat,prices,price_hat,mode);rows.append(r);cache[('q42',mode,0,.95)]=ds;print('42',mode,r['total_cost'])
 for name,mode,tail in [('deterministic','joint_block',False),('robust','joint_block',False),('DRO_tail','joint_block',True)]:
  r,ds=q42_backtest(name,net,net_hat,prices,price_hat,mode,tail=tail);rows.append(r);cache[('q42',name,0,.95)]=ds;print('42',name,r['total_cost'])
 for alpha in [.9,.95,.99]:
  for lam in [.02,.1,.3]:
   r,ds=q42_backtest('CVaR',net,net_hat,prices,price_hat,'joint_block',lam,alpha);rows.append(r);cache[('q42','CVaR',lam,alpha)]=ds;print('42 cvar',alpha,lam,r['total_cost'])
 # Q4-3 principal methods and information decomposition.
 r,ds=q43_backtest('deterministic_dynamic_price',L,P,F,Lhat,prices_all,np.vstack([prices_all[:START],price_hat]),deterministic=True);rows.append(r);cache[('q43','det')]=ds;print('43 det',r['total_cost'])
 for mode in ['independent','joint_block']:
  r,ds=q43_backtest('stochastic_dynamic_price',L,P,F,Lhat,prices_all,np.vstack([prices_all[:START],price_hat]),mode=mode);rows.append(r);cache[('q43',mode,0,.95)]=ds;print('43',mode,r['total_cost'])
 for lam in [.02,.1,.3]:
  r,ds=q43_backtest('risk_aware_MPC',L,P,F,Lhat,prices_all,np.vstack([prices_all[:START],price_hat]),risk_lambda=lam,alpha=.95);rows.append(r);cache[('q43','risk',lam,.95)]=ds;print('43 risk',lam,r['total_cost'])
 for pvup,prup in [(False,False),(True,False),(False,True),(True,True)]:
  r,ds=q43_backtest('information_ablation',L,P,F,Lhat,prices_all,np.vstack([prices_all[:START],price_hat]),pv_updates=pvup,price_updates=prup);rows.append(r);cache[('info',pvup,prup)]=ds;print('info',pvup,prup,r['total_cost'])
 # Final selection: minimum realized mean cost; CVaR only if q95 improves materially with <=2% mean penalty.
 q42base=min([x for x in rows if x['problem']=='Q4-2' and x['name']!='CVaR'],key=lambda x:x['total_cost']);cands=[x for x in rows if x['problem']=='Q4-2' and x['name']=='CVaR' and x['daily_q95']<.98*q42base['daily_q95'] and x['total_cost']<1.02*q42base['total_cost']];q42win=min(cands,key=lambda x:x['daily_q95']) if cands else q42base
 if q42win['name']=='CVaR':d42=cache[('q42','CVaR',q42win['risk_lambda'],q42win['alpha'])]
 elif q42win['name']=='joint_structure':d42=cache[('q42',q42win['scenario'],0,.95)]
 else:d42=cache[('q42',q42win['name'],0,.95)]
 q43base=min([x for x in rows if x['problem']=='Q4-3' and x['name'] in ('deterministic_dynamic_price','stochastic_dynamic_price')],key=lambda x:x['total_cost']);cands=[x for x in rows if x['problem']=='Q4-3' and x['name']=='risk_aware_MPC' and x['daily_q95']<.98*q43base['daily_q95'] and x['total_cost']<1.02*q43base['total_cost']];q43win=min(cands,key=lambda x:x['daily_q95']) if cands else q43base
 if q43win['name']=='deterministic_dynamic_price':d43=cache[('q43','det')]
 elif q43win['name']=='risk_aware_MPC':d43=cache[('q43','risk',q43win['risk_lambda'],q43win['alpha'])]
 else:d43=cache[('q43',q43win['scenario'],0,.95)]
 save_result42(dates,d42);save_result43(dates,d43)
 fields=[]
 for x in rows:
  for k in x:
   if k not in fields:fields.append(k)
 with (OUT/'economic_experiments.csv').open('w',newline='',encoding='utf-8-sig') as f:w=csv.DictWriter(f,fieldnames=fields);w.writeheader();w.writerows(rows)
 summary={'price_model':'HistGB_XGB_fallback','q42_winner':q42win,'q43_winner':q43win,'information':[{k:x[k] for k in ('pv_updates','price_updates','total_cost','emergency_cost')} for x in rows if x['name']=='information_ablation']}
 (OUT/'summary.json').write_text(json.dumps(summary,ensure_ascii=False,indent=2,default=float),encoding='utf-8')
 np.savez_compressed(OUT/'final_q42.npz',**{k:np.array([x[k] for x in d42]) for k in ['g','c','d','soc','e']},cost=np.array([x['total_cost'] for x in d42]))
 np.savez_compressed(OUT/'final_q43.npz',**{k:np.array([x[k] for x in d43]) for k in ['g0','gfinal','c','d','soc','e']},cost=np.array([x['total_cost'] for x in d43]))
 print(json.dumps(summary,ensure_ascii=False,indent=2,default=float))
if __name__=='__main__':main()

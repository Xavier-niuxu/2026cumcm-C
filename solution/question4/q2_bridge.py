# -*- coding: utf-8 -*-
"""替代参考解中缺失的 ``q2/forecast_benchmark``：提供 ``load_data()`` 与 ``fourier()``。

参考解的 ``q4_models.py`` 从 ``../q2/forecast_benchmark.py`` 导入 ``load_data`` 与 ``fourier``，
但该模块不在交付包里。这里用本仓库问题 2 的傅里叶模型补齐（与问题 3 的处理方式一致）：

* ``load_data()``   -> (小区负载 L, 实际光伏 P, 净负荷 N = L - P)，均为 365×144
* ``fourier(X, d)`` -> 对矩阵 X（负荷或净负荷）做**严格因果**的日前傅里叶预报，返回第 d 天 144 维

问题 2 的模型建模的是净负荷，把第二个参数（光伏）置 0 即可对任意单序列建模，
因此 L 与 N 用同一套模型分别拟合。
"""

from __future__ import annotations

import sys
from pathlib import Path

import numpy as np
from openpyxl import load_workbook

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
ATT = ROOT / "附件"
Q2_DIR = HERE.parent / "question2"
if str(Q2_DIR) not in sys.path:
    sys.path.append(str(Q2_DIR))

from pred.fourier import FourierPredictor  # noqa: E402

_CACHE: dict = {}


def _matrix(path: Path, sheet_index: int = 0):
    w = load_workbook(path, read_only=True, data_only=True)
    s = w.worksheets[sheet_index]
    dates, rows = [], []
    for r in s.iter_rows(min_row=2, values_only=True):
        dates.append(r[0])
        rows.append([float(x) for x in r[1:]])
    return np.asarray(dates), np.asarray(rows, float)


def load_data():
    """返回 (负荷 L, 光伏 P, 净负荷 L-P)。"""
    dates, L = _matrix(ATT / "附件2.xlsx", 0)
    _, P = _matrix(ATT / "附件2.xlsx", 1)
    return L, P, L - P


def _model_for(key: str, series: np.ndarray):
    """按序列名缓存一个因果傅里叶模型（光伏置 0，等价于对该序列本身建模）。"""
    if key not in _CACHE:
        w = load_workbook(ATT / "附件2.xlsx", read_only=True, data_only=True)
        dates = np.array([r[0] for r in w.worksheets[0].iter_rows(min_row=2, values_only=True)])
        _CACHE[key] = (dates, FourierPredictor(dates, series, dates, np.zeros_like(series)))
    return _CACHE[key]


def fourier(X: np.ndarray, d: int, *, key: str | None = None) -> np.ndarray:
    """第 d 天的日前预报（144 维），只用 d 之前的数据拟合。"""
    key = key or f"series-{id(X)}"
    dates, model = _model_for(key, np.asarray(X, float))
    pred, _ = model.predict(str(dates[d])[:10])
    return np.asarray(pred, float)
